import io
import os
import zipfile
import base64
import logging
from typing import Optional
import httpx
from fastapi import HTTPException, status
import pypdf
import docx

from app.config import settings

logger = logging.getLogger("decisionlens.file_parser")

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}


def extract_scanned_pdf_ocr(file_bytes: bytes) -> Optional[str]:
    """
    Fallback OCR for scanned PDFs without selectable digital text.
    Uses Google Gemini Flash native PDF understanding or Azure OpenAI Vision.
    """
    # 1. Tier 0: Google Gemini multimodal PDF transcription (Fast & native PDF input)
    gemini_key = (settings.gemini_api_key or settings.google_api_key).strip()
    if gemini_key:
        b64_pdf = base64.b64encode(file_bytes).decode("utf-8")
        models_to_try = ["gemini-2.5-flash", "gemini-3.5-flash", "gemini-3.1-flash-lite", "gemini-flash-latest"]
        for model_name in models_to_try:
            try:
                url = f"https://generativelanguage.googleapis.com/v1beta/models/{model_name}:generateContent?key={gemini_key}"
                payload = {
                    "contents": [
                        {
                            "parts": [
                                {
                                    "inline_data": {
                                        "mime_type": "application/pdf",
                                        "data": b64_pdf,
                                    }
                                },
                                {
                                    "text": (
                                        "Transcribe and extract all readable text, headings, numbers, and tables "
                                        "from this scanned document verbatim. Return only the extracted text formatted cleanly in markdown."
                                    )
                                },
                            ]
                        }
                    ],
                    "generationConfig": {
                        "temperature": 0.1,
                        "maxOutputTokens": 4096,
                    },
                }
                with httpx.Client(timeout=40.0) as client:
                    resp = client.post(url, json=payload)
                    if resp.status_code == 200:
                        data = resp.json()
                        candidates = data.get("candidates", [])
                        if candidates:
                            parts = candidates[0].get("content", {}).get("parts", [])
                            extracted = "".join(p.get("text", "") for p in parts).strip()
                            if extracted:
                                logger.info("Successfully extracted text from scanned PDF via Gemini %s", model_name)
                                return f"[Scanned Document OCR Text]\n{extracted}"
            except Exception as e:
                logger.warning("Gemini OCR attempt failed with model %s: %s", model_name, e)

    # 2. Tier 1: Azure OpenAI Vision (extract page images and transcribe)
    try:
        from app.services.ai import get_llm_client_and_model
        client, model, provider = get_llm_client_and_model()
        if client and provider in {"azure", "openai"}:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            page_texts = []
            for i, page in enumerate(reader.pages[:8]):  # Process up to 8 pages
                images = list(page.images)
                if not images:
                    continue
                largest_img = max(images, key=lambda img: len(img.data))
                img_b64 = base64.b64encode(largest_img.data).decode("utf-8")
                ext_name = largest_img.name.split(".")[-1].lower()
                mime = "image/png" if ext_name == "png" else "image/jpeg"

                resp = client.chat.completions.create(
                    model=model,
                    messages=[
                        {
                            "role": "user",
                            "content": [
                                {
                                    "type": "text",
                                    "text": "Transcribe all readable text from this scanned document page exactly as written.",
                                },
                                {
                                    "type": "image_url",
                                    "image_url": {"url": f"data:{mime};base64,{img_b64}"},
                                },
                            ],
                        }
                    ],
                    max_tokens=2048,
                    temperature=0.1,
                )
                txt = resp.choices[0].message.content or ""
                if txt.strip():
                    page_texts.append(f"--- Page {i + 1} (OCR) ---\n{txt.strip()}")

            if page_texts:
                logger.info("Successfully extracted text from scanned PDF via Azure OpenAI Vision (%d pages)", len(page_texts))
                return "\n\n".join(page_texts)
    except Exception as e:
        logger.warning("Azure OpenAI Vision OCR attempt failed: %s", e)

    return None


def get_file_extension(filename: str) -> str:
    """Return normalized lowercase file extension including the dot."""
    _, ext = os.path.splitext(filename)
    return ext.lower()


def extract_text_from_file(file_bytes: bytes, filename: str, mime_type: Optional[str] = None) -> str:
    """
    Extract readable text from PDF, DOCX, TXT, and Markdown files.
    Rejects empty files, unsupported extensions, and corrupted binaries.
    Automatically applies AI Vision OCR for scanned PDFs without selectable text.
    """
    if not file_bytes or len(file_bytes) == 0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Uploaded file is empty (0 bytes).",
        )

    ext = get_file_extension(filename)
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file format '{ext}'. Supported formats: .pdf, .docx, .txt, .md",
        )

    # 1. Plain Text and Markdown
    if ext in {".txt", ".md", ".markdown"}:
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = file_bytes.decode("latin-1")
            except Exception:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Failed to decode text file. Ensure it is encoded in UTF-8 or standard ASCII.",
                )

        cleaned_text = text.strip()
        if not cleaned_text:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="File is empty or contains only whitespace.",
            )
        return cleaned_text

    # 2. Microsoft Word Document (.docx)
    elif ext == ".docx":
        try:
            stream = io.BytesIO(file_bytes)
            doc = docx.Document(stream)
            extracted_paragraphs = []

            for p in doc.paragraphs:
                p_text = p.text.strip()
                if p_text:
                    extracted_paragraphs.append(p_text)

            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        extracted_paragraphs.append(" | ".join(row_cells))

            full_text = "\n\n".join(extracted_paragraphs).strip()
            if not full_text:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Word document (.docx) contains no extractable text.",
                )
            return full_text

        except (zipfile.BadZipFile, docx.opc.exceptions.PackageNotFoundError, Exception) as exc:
            if isinstance(exc, HTTPException):
                raise exc
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Malformed or corrupted DOCX file: {str(exc)}",
            )

    # 3. PDF Document (.pdf)
    elif ext == ".pdf":
        try:
            stream = io.BytesIO(file_bytes)
            reader = pypdf.PdfReader(stream)

            if len(reader.pages) == 0:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="PDF contains no pages.",
                )

            extracted_pages = []
            for i, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    extracted_pages.append(f"--- Page {i + 1} ---\n{page_text.strip()}")

            full_text = "\n\n".join(extracted_pages).strip()

            if not full_text:
                # Scanned PDF with zero selectable text: Fallback to AI Vision OCR
                logger.info("PDF '%s' has 0 selectable characters. Initiating AI Vision OCR...", filename)
                ocr_result = extract_scanned_pdf_ocr(file_bytes)
                if ocr_result and ocr_result.strip():
                    return ocr_result.strip()

                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail="Scanned PDF detected without selectable text. AI OCR was unable to read the pages. Please upload a clearer copy or paste the text directly.",
                )

            return full_text

        except Exception as exc:
            if isinstance(exc, HTTPException):
                raise exc
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Malformed or unreadable PDF document: {str(exc)}",
            )

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail=f"Unsupported file format '{ext}'.",
    )
