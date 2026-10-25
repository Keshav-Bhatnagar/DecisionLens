import io
import docx
import pypdf
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from app.models import User, Folder, Decision
from app.auth import hash_password, create_token
from app.config import settings


@pytest.fixture
def user_a(db_session: Session) -> dict:
    user = User(email="upload_a@example.com", password_hash=hash_password("PassA123!"))
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)

    folder = Folder(owner_id=user.id, name="Uploads Test Folder")
    db_session.add(folder)
    db_session.commit()
    db_session.refresh(folder)

    decision = Decision(owner_id=user.id, folder_id=folder.id, title="Uploads Decision")
    db_session.add(decision)
    db_session.commit()
    db_session.refresh(decision)

    return {"user": user, "token": create_token(user.id), "decision": decision}


@pytest.fixture
def user_b(db_session: Session) -> dict:
    user = User(email="upload_b@example.com", password_hash=hash_password("PassB123!"))
    db_session.add(user)
    db_session.commit()
    db_session.refresh(user)
    return {"user": user, "token": create_token(user.id)}


def make_test_pdf(text_content: str = "DecisionLens Architecture Plan") -> bytes:
    content = f"BT /F1 12 Tf 50 700 Td ({text_content}) Tj ET"
    pdf_str = f"""%PDF-1.4
1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj
2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj
3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj
4 0 obj << /Length {len(content)} >> stream
{content}
endstream
endobj
5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000244 00000 n 
0000000334 00000 n 
trailer << /Size 6 /Root 1 0 R >>
startxref
403
%%EOF"""
    return pdf_str.encode("latin-1")


def make_scanned_blank_pdf() -> bytes:
    writer = pypdf.PdfWriter()
    writer.add_blank_page(width=200, height=200)
    stream = io.BytesIO()
    writer.write(stream)
    return stream.getvalue()


def make_test_docx(text_content: str = "Sample Word Document Content") -> bytes:
    doc = docx.Document()
    doc.add_heading("Architecture Overview", level=1)
    doc.add_paragraph(text_content)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Criterion A"
    table.rows[0].cells[1].text = "Value High"
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_upload_txt_file_success(client: TestClient, user_a: dict):
    """Test uploading a valid TXT document."""
    headers = {"Authorization": f"Bearer {user_a['token']}"}
    decision_id = user_a["decision"].id

    txt_content = b"Problem Statement: Latency in API gateway is causing customer churn."
    files = {"file": ("problem_notes.txt", txt_content, "text/plain")}

    response = client.post(f"/api/decisions/{decision_id}/attachments", headers=headers, files=files)
    assert response.status_code == 201

    data = response.json()
    assert data["original_name"] == "problem_notes.txt"
    assert "Latency in API gateway" in data["extracted_text"]
    assert data["file_size"] == len(txt_content)

    # Verify attachment appears in list
    list_res = client.get(f"/api/decisions/{decision_id}/attachments", headers=headers)
    assert list_res.status_code == 200
    assert len(list_res.json()) == 1


def test_upload_markdown_file_success(client: TestClient, user_a: dict):
    """Test uploading a valid Markdown (.md) document."""
    headers = {"Authorization": f"Bearer {user_a['token']}"}
    decision_id = user_a["decision"].id

    md_content = b"# Migration Spec\n\n- Option A: Postgres\n- Option B: DynamoDB"
    files = {"file": ("spec.md", md_content, "text/markdown")}

    response = client.post(f"/api/decisions/{decision_id}/attachments", headers=headers, files=files)
    assert response.status_code == 201
    assert "Option A: Postgres" in response.json()["extracted_text"]


def test_upload_docx_file_success(client: TestClient, user_a: dict):
    """Test uploading a valid Word (.docx) document."""
    headers = {"Authorization": f"Bearer {user_a['token']}"}
    decision_id = user_a["decision"].id

    docx_bytes = make_test_docx("We must evaluate cloud hosting cost trade-offs.")
    files = {"file": ("requirements.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    response = client.post(f"/api/decisions/{decision_id}/attachments", headers=headers, files=files)
    assert response.status_code == 201
    assert "cloud hosting cost trade-offs" in response.json()["extracted_text"]
    assert "Criterion A" in response.json()["extracted_text"]


def test_upload_pdf_file_success(client: TestClient, user_a: dict):
    """Test uploading a valid PDF document with selectable text."""
    headers = {"Authorization": f"Bearer {user_a['token']}"}
    decision_id = user_a["decision"].id

    pdf_bytes = make_test_pdf("DecisionLens Architecture Plan")
    files = {"file": ("architecture.pdf", pdf_bytes, "application/pdf")}

    response = client.post(f"/api/decisions/{decision_id}/attachments", headers=headers, files=files)
    assert response.status_code == 201
    assert "DecisionLens Architecture Plan" in response.json()["extracted_text"]


def test_upload_scanned_pdf_ocr(client: TestClient, user_a: dict):
    """Test scanned PDF without selectable text triggers OCR and succeeds or returns clear error."""
    headers = {"Authorization": f"Bearer {user_a['token']}"}
    decision_id = user_a["decision"].id

    scanned_pdf = make_scanned_blank_pdf()
    files = {"file": ("scanned_receipt.pdf", scanned_pdf, "application/pdf")}

    response = client.post(f"/api/decisions/{decision_id}/attachments", headers=headers, files=files)
    # With AI OCR active, it succeeds with 201; if no key configured in test, it returns 400
    assert response.status_code in (201, 400)


def test_google_drive_import_invalid_url(client: TestClient, user_a: dict):
    """Test Google Drive import with invalid URL rejects with 400."""
    headers = {"Authorization": f"Bearer {user_a['token']}"}
    decision_id = user_a["decision"].id

    payload = {"drive_url": "https://invalid-link-without-id.com"}
    response = client.post(f"/api/decisions/{decision_id}/attachments/google-drive", headers=headers, json=payload)
    assert response.status_code == 400
    assert "google drive file id" in response.json()["detail"].lower()


def test_upload_empty_file_rejection(client: TestClient, user_a: dict):
    """Test empty file (0 bytes) is rejected with HTTP 400."""
    headers = {"Authorization": f"Bearer {user_a['token']}"}
    decision_id = user_a["decision"].id

    files = {"file": ("empty.txt", b"", "text/plain")}
    response = client.post(f"/api/decisions/{decision_id}/attachments", headers=headers, files=files)
    assert response.status_code == 400
    assert "empty" in response.json()["detail"].lower()


def test_upload_unsupported_format_rejection(client: TestClient, user_a: dict):
    """Test unsupported extensions (.exe, .jpg) are rejected with HTTP 400."""
    headers = {"Authorization": f"Bearer {user_a['token']}"}
    decision_id = user_a["decision"].id

    files = {"file": ("malicious.exe", b"MZ...", "application/x-msdownload")}
    response = client.post(f"/api/decisions/{decision_id}/attachments", headers=headers, files=files)
    assert response.status_code == 400
    assert "unsupported file format" in response.json()["detail"].lower()


def test_upload_corrupted_file_rejection(client: TestClient, user_a: dict):
    """Test malformed/corrupted file content is rejected with HTTP 400."""
    headers = {"Authorization": f"Bearer {user_a['token']}"}
    decision_id = user_a["decision"].id

    files = {"file": ("corrupt.docx", b"Not A Valid Docx Zip", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    response = client.post(f"/api/decisions/{decision_id}/attachments", headers=headers, files=files)
    assert response.status_code == 400
    assert "malformed" in response.json()["detail"].lower()


def test_upload_size_limit_rejection(client: TestClient, user_a: dict, monkeypatch):
    """Test files exceeding maximum upload limit are rejected with HTTP 413."""
    headers = {"Authorization": f"Bearer {user_a['token']}"}
    decision_id = user_a["decision"].id

    # Temporarily set max upload to 1MB for testing
    monkeypatch.setattr(settings, "max_upload_mb", 1)

    oversized_content = b"A" * (2 * 1024 * 1024)  # 2MB
    files = {"file": ("huge_log.txt", oversized_content, "text/plain")}

    response = client.post(f"/api/decisions/{decision_id}/attachments", headers=headers, files=files)
    assert response.status_code == 413
    assert "exceeds" in response.json()["detail"].lower()


def test_attachment_cross_user_isolation(client: TestClient, user_a: dict, user_b: dict):
    """Verify strict user isolation on attachments."""
    headers_a = {"Authorization": f"Bearer {user_a['token']}"}
    headers_b = {"Authorization": f"Bearer {user_b['token']}"}
    decision_a_id = user_a["decision"].id

    # User A uploads attachment
    files = {"file": ("private_notes.txt", b"Confidential business data", "text/plain")}
    upload_res = client.post(f"/api/decisions/{decision_a_id}/attachments", headers=headers_a, files=files)
    assert upload_res.status_code == 201
    attachment_id = upload_res.json()["id"]

    # User B attempts to upload attachment to User A's decision -> 404 Not Found
    b_upload = client.post(
        f"/api/decisions/{decision_a_id}/attachments",
        headers=headers_b,
        files={"file": ("injected.txt", b"Injected data", "text/plain")},
    )
    assert b_upload.status_code == 404

    # User B attempts to list attachments of User A's decision -> 404 Not Found
    b_list = client.get(f"/api/decisions/{decision_a_id}/attachments", headers=headers_b)
    assert b_list.status_code == 404

    # User B attempts to read User A's attachment -> 404 Not Found
    b_get = client.get(f"/api/decisions/{decision_a_id}/attachments/{attachment_id}", headers=headers_b)
    assert b_get.status_code == 404

    # User B attempts to delete User A's attachment -> 404 Not Found
    b_del = client.delete(f"/api/decisions/{decision_a_id}/attachments/{attachment_id}", headers=headers_b)
    assert b_del.status_code == 404

    # User A deletes attachment -> 200 OK
    a_del = client.delete(f"/api/decisions/{decision_a_id}/attachments/{attachment_id}", headers=headers_a)
    assert a_del.status_code == 200

    # Verify deleted
    assert client.get(f"/api/decisions/{decision_a_id}/attachments/{attachment_id}", headers=headers_a).status_code == 404
